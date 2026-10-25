"""运行方读用户数据（SDD 0027 §3.5，knowbot-data/v1）。

knowlab 这一道核、抽文字与按格子读（0030 §3.1）。
"""

from __future__ import annotations

import asyncio
import datetime as dt
import hashlib
import io
import json
import os
import threading
import uuid
from dataclasses import replace
from pathlib import Path

import httpx
import pytest

from app.application.errors.exceptions import NotFoundError, ValidationError
from app.application.ports.knowbots import DatasetRecord, ExpertRecord, TestRecord
from app.application.services.knowbot_data_read import KnowbotDataReader
from app.domain.knowbots.data_read import (
    PAGE_BYTES,
    CellsRequestV1,
    Extracted,
    ListRequestV1,
    QuotesRequestV1,
    ReadRequestV1,
    Sheet,
    TableRequestV1,
    TableUnreadable,
    find_quote,
    normalize_quote,
    paginate,
    pdf_pages,
    table_page,
)
from app.domain.knowbots.data_schema import DOCUMENTS, render
from app.domain.knowbots.datasets import XLSX
from app.infrastructure.documents import table_extract, text_extract
from app.infrastructure.documents.table_extract import (
    CachedTableExtractor,
    _sheet,
    _xls_value,
    extract_tables,
    number_of,
)
from app.infrastructure.documents.text_extract import (
    NO_TEXT,
    UNREADABLE,
    CachedTextExtractor,
    extract_text,
)
from app.infrastructure.security.relation_auth import (
    ServiceCaller,
    require_read_service,
)
from app.interfaces.endpoints import knowbot_data_routes
from app.main import app
from tests.knowbot_fakes import OWNER, MemoryDataStore, MemoryKnowlab, manifest

KNOWBOT = "0b9b8c52-3f0e-4c34-9d0b-2b7f4d3c1a10"
INTERNAL = "/api/internal/v1/knowlab/knowbot-data"
CONTRACT_DIR = Path(__file__).resolve().parents[3] / "contracts/knowbot-data/v1"


class World:
    """一个用户、两份数据（一份绑定、一份没绑）、一版上架了的专家、一位在测试的准专家。"""

    def __init__(self) -> None:
        self.repo = MemoryKnowlab()
        self.store = MemoryDataStore()
        self.reader = KnowbotDataReader(
            self.repo, self.store, CachedTextExtractor(), CachedTableExtractor()
        )

    async def setup(self) -> None:
        self.bound = await self._dataset(
            "年报.csv", "text/csv; charset=utf-8", "公司,营收\nA,100\n"
        )
        self.loose = await self._dataset(
            "别的.txt", "text/plain; charset=utf-8", "不该读到"
        )
        binding = {
            "dataset_id": str(self.bound.id),
            "version": "1",
            "sha256": self.bound.sha256,
        }
        self.request = await self.repo.create_request(
            owner_subject=OWNER, title="t", runtime="knowbench", idea={}, cases=[]
        )
        doc = manifest("1", data_bindings=[binding])
        self.digest = "sha256:" + "a" * 64
        self.delivered = await self.repo.create_draft(
            request_id=self.request.id,
            request_version=1,
            revision=1,
            version="1",
            manifest=doc,
        )
        await self.repo.update_draft(self.delivered.id, status="passed")
        now = self.repo._now()
        await self.repo.create_expert(
            ExpertRecord(
                id=uuid.uuid4(),
                request_id=self.request.id,
                version=1,
                draft_id=self.delivered.id,
                created_at=now,
                listed_at=now,
                knowmart_knowbot_id=KNOWBOT,
                digest=self.digest,
            )
        )
        self.testing = await self.repo.create_draft(
            request_id=self.request.id,
            request_version=1,
            revision=2,
            version="2",
            manifest=manifest("2", data_bindings=[binding]),
        )
        await self.repo.update_draft(self.testing.id, status="testing")
        self.test = await self.repo.create_test(
            TestRecord(
                id=uuid.uuid4(),
                request_id=self.request.id,
                number=1,
                draft_id=self.testing.id,
                status="running",
                run_id="knowlab-run-0002",
                started_at=now,
            )
        )

    async def _dataset(
        self, name: str, content_type: str, text: str | bytes
    ) -> DatasetRecord:
        data = text.encode("utf-8") if isinstance(text, str) else text
        record_id = uuid.uuid4()
        key = f"users/x/{record_id}"
        version_id, sha = await self.store.put(
            object_key=key, data=data, content_type=content_type
        )
        return await self.repo.create_dataset(
            DatasetRecord(
                id=record_id,
                owner_subject=OWNER,
                name=name,
                version="1",
                content_type=content_type,
                size_bytes=len(data),
                sha256=sha,
                object_key=key,
                version_id=version_id,
            )
        )

    def knowbot(self, **overrides: str) -> dict[str, str]:
        purpose = {
            "kind": "knowbot",
            "owner_subject": OWNER,
            "knowbot_id": KNOWBOT,
            "version": "1",
            "digest": self.digest,
        }
        return purpose | overrides

    def listing(self, purpose: dict[str, str]) -> ListRequestV1:
        return ListRequestV1.model_validate(
            {"contract": "knowbot-data/v1", "purpose": purpose}
        )

    def reading(
        self, purpose: dict[str, str], dataset: str, page: int = 1
    ) -> ReadRequestV1:
        return ReadRequestV1.model_validate(
            {
                "contract": "knowbot-data/v1",
                "purpose": purpose,
                "dataset_id": dataset,
                "page": page,
            }
        )

    def table(
        self, purpose: dict[str, str], dataset: str, sheet: str | None = None
    ) -> TableRequestV1:
        return TableRequestV1.model_validate(
            {
                "contract": "knowbot-data/v1",
                "purpose": purpose,
                "dataset_id": dataset,
                "sheet": sheet,
            }
        )

    def cells(
        self, purpose: dict[str, str], *refs: tuple[str, str, str, str]
    ) -> CellsRequestV1:
        return CellsRequestV1.model_validate(
            {
                "contract": "knowbot-data/v1",
                "purpose": purpose,
                "cells": [
                    {"dataset_id": d, "sheet": s, "row": r, "column": c}
                    for d, s, r, c in refs
                ],
            }
        )


@pytest.fixture
async def world() -> World:
    w = World()
    await w.setup()
    return w


async def test_a_delivered_knowbot_lists_and_reads_only_its_bound_data(
    world: World,
) -> None:
    listed = await world.reader.list("knowbench", world.listing(world.knowbot()))
    assert [d.dataset_id for d in listed.datasets] == [str(world.bound.id)]
    read = await world.reader.read(
        "knowbench", world.reading(world.knowbot(), str(world.bound.id))
    )
    assert read.text == "公司,营收\nA,100\n"
    assert (read.page, read.pages, read.truncated, read.note) == (1, 1, False, None)
    assert read.citation.sha256 == world.bound.sha256 and read.citation.version == "1"


async def test_a_running_test_reads_the_same_way(world: World) -> None:
    purpose = {"kind": "test", "run_id": "knowlab-run-0002"}
    read = await world.reader.read(
        "knowbench", world.reading(purpose, str(world.bound.id))
    )
    assert read.name == "年报.csv"


async def test_a_test_reads_here_only_what_is_not_in_knowbase(world: World) -> None:
    """测试签了运行授权、或这一份已搬进 knowbase：只能凭授权读（Astra F02）。"""
    purpose = {"kind": "test", "run_id": "knowlab-run-0002"}
    # 搬过来、旧对象还在（本人还没补上测试同意时开的测试）：照旧经这里读（Fable P2-4）
    world.repo.datasets[world.bound.id] = replace(world.bound, item_id="item-1")
    read = await world.reader.read(
        "knowbench", world.reading(purpose, str(world.bound.id))
    )
    assert read.name == "年报.csv"
    # 上架的第 1 版照旧经这里读，搬没搬都一样（SDD 0042 第七节）
    read = await world.reader.read(
        "knowbench", world.reading(world.knowbot(), str(world.bound.id))
    )
    assert read.name == "年报.csv"
    # 只在 knowbase 里的：只能凭授权读
    world.repo.datasets[world.bound.id] = replace(
        world.bound, item_id="item-1", object_key=None, version_id=None
    )
    with pytest.raises(NotFoundError):
        await world.reader.read(
            "knowbench", world.reading(purpose, str(world.bound.id))
        )
    world.repo.datasets[world.bound.id] = world.bound
    await world.repo.update_test(world.test.id, grant_id="grant-1")
    with pytest.raises(NotFoundError):
        await world.reader.list("knowbench", world.listing(purpose))


@pytest.mark.parametrize(
    "change",
    [
        {"owner_subject": "someone-else"},
        {"digest": "sha256:" + "b" * 64},
        {"version": "2"},
        {"knowbot_id": "11111111-1111-4111-8111-111111111111"},
    ],
)
async def test_anything_off_from_the_delivery_is_not_found(
    world: World, change: dict[str, str]
) -> None:
    with pytest.raises(NotFoundError) as caught:
        await world.reader.list("knowbench", world.listing(world.knowbot(**change)))
    assert caught.value.code == "knowbot_data_not_found"


async def test_another_runtime_cannot_read(world: World) -> None:
    with pytest.raises(NotFoundError):
        await world.reader.list("other-runtime", world.listing(world.knowbot()))


async def test_unbound_or_deleted_or_changed_data_is_not_found(world: World) -> None:
    with pytest.raises(NotFoundError):
        await world.reader.read(
            "knowbench", world.reading(world.knowbot(), str(world.loose.id))
        )
    world.repo.datasets[world.bound.id] = replace(world.bound, sha256="c" * 64)
    with pytest.raises(NotFoundError):
        await world.reader.read(
            "knowbench", world.reading(world.knowbot(), str(world.bound.id))
        )
    world.repo.datasets[world.bound.id] = replace(
        world.bound, deleted_at=world.repo._now()
    )
    with pytest.raises(NotFoundError):
        await world.reader.read(
            "knowbench", world.reading(world.knowbot(), str(world.bound.id))
        )
    assert (
        await world.reader.list("knowbench", world.listing(world.knowbot()))
    ).datasets == ()


async def test_a_finished_or_other_test_cannot_read(world: World) -> None:
    for run_id in ("knowlab-run-9999", "knowlab-run-0002"):
        if run_id == "knowlab-run-0002":
            await world.repo.update_test(world.test.id, status="passed")
        with pytest.raises(NotFoundError):
            await world.reader.list(
                "knowbench", world.listing({"kind": "test", "run_id": run_id})
            )


async def test_pages_beyond_the_end_say_so(world: World) -> None:
    read = await world.reader.read(
        "knowbench", world.reading(world.knowbot(), str(world.bound.id), page=3)
    )
    assert (read.page, read.pages, read.text) == (3, 1, "")
    assert read.note == "There are only 1 pages."


async def test_a_bound_csv_reads_as_a_table_and_by_cell(world: World) -> None:
    dataset = str(world.bound.id)
    table = await world.reader.table("knowbench", world.table(world.knowbot(), dataset))
    assert (table.sheets, table.sheet, table.columns) == (("data",), "data", ("营收",))
    assert [(r.label, r.values) for r in table.rows] == [("A", {"营收": "100"})]
    assert table.citation.sha256 == world.bound.sha256
    got = await world.reader.cells(
        "knowbench",
        world.cells(
            world.knowbot(),
            (dataset, "data", "A", "营收"),
            (dataset, "data", "B", "营收"),
            (dataset, "data", "A", "利润"),
            (dataset, "表2", "A", "营收"),
        ),
    )
    assert [(c.found, c.raw, c.number) for c in got.cells] == [
        (True, "100", "100"),
        (False, None, None),
        (False, None, None),
        (False, None, None),
    ]
    assert {c.version for c in got.cells} == {"1"}


async def test_tables_and_cells_keep_the_same_two_checks(world: World) -> None:
    loose = str(world.loose.id)
    with pytest.raises(NotFoundError):
        await world.reader.table("knowbench", world.table(world.knowbot(), loose))
    bound = str(world.bound.id)
    with pytest.raises(NotFoundError):
        await world.reader.cells(
            "knowbench",
            world.cells(
                world.knowbot(), (bound, "data", "A", "营收"), (loose, "x", "y", "z")
            ),
        )
    with pytest.raises(NotFoundError):
        await world.reader.cells(
            "otherapp", world.cells(world.knowbot(), (bound, "data", "A", "营收"))
        )
    test = {"kind": "test", "run_id": "knowlab-run-0002"}
    got = await world.reader.cells(
        "knowbench", world.cells(test, (bound, "data", "A", "营收"))
    )
    assert got.cells[0].number == "100"


async def test_text_data_or_a_missing_sheet_is_not_a_table(world: World) -> None:
    text = await world._dataset("说明.txt", "text/plain; charset=utf-8", "就是字")
    binding = {"dataset_id": str(text.id), "version": "1", "sha256": text.sha256}
    draft = await world.repo.get_draft(world.testing.id)
    assert draft is not None
    doc = dict(draft.manifest)
    doc["data_bindings"] = [*doc["data_bindings"], binding]
    await world.repo.update_draft(world.testing.id, manifest=doc)
    test = {"kind": "test", "run_id": "knowlab-run-0002"}
    with pytest.raises(ValidationError) as caught:
        await world.reader.table("knowbench", world.table(test, str(text.id)))
    assert caught.value.code == "knowbot_data_not_a_table"
    with pytest.raises(ValidationError) as caught:
        await world.reader.table(
            "knowbench", world.table(test, str(world.bound.id), "表2")
        )
    assert caught.value.code == "knowbot_data_no_sheet"


def test_paginate_splits_at_the_page_size() -> None:
    text = "x" * (PAGE_BYTES * 2 + 5)  # 一行就超了：按字切开
    assert paginate(text, 1) == ("x" * PAGE_BYTES, 3)
    assert paginate(text, 3) == ("x" * 5, 3)
    assert paginate("", 1) == ("", 1)


def test_a_chinese_annual_report_is_paged_by_bytes_and_keeps_its_pages_whole() -> None:
    """中文年报按字节切（待办账 137）。

    每页都在模型看得到的上限里，在「## Page N」前面切，接起来就是原文。
    """
    printed = [
        f"## Page {n}\n"
        + "".join(f"营业收入{n}-{i}，170,899,152,276.34 元\n" for i in range(400))
        for n in range(1, 13)
    ]
    text = "".join(printed)
    first, pages = paginate(text, 1)
    parts = [paginate(text, p)[0] for p in range(1, pages + 1)]
    assert "".join(parts) == text
    assert all(len(part.encode("utf-8")) <= PAGE_BYTES for part in parts)
    assert all(part.startswith("## Page ") for part in parts)  # 没把年报的一页拆开
    assert pages > 1 and paginate(text, pages + 1) == ("", pages)


def test_requests_outside_the_contract_are_rejected() -> None:
    with pytest.raises(ValueError):
        ListRequestV1.model_validate(
            {"contract": "knowbot-data/v1", "purpose": {"kind": "all"}}
        )
    with pytest.raises(ValueError):
        ReadRequestV1.model_validate(
            {
                "contract": "knowbot-data/v1",
                "purpose": {"kind": "test", "run_id": "knowlab-run-1"},
                "dataset_id": "../etc/passwd",
            }
        )


# ---------- 抽文字 ----------


def _xlsx() -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "利润表"
    sheet.append(["年份", "营收", None])
    sheet.append([2025, 100.0, None])
    book.create_sheet("说明").append(["单位：万元"])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _pdf(text: str) -> bytes:
    """最小的一页 PDF，内容是一行 ASCII 字。"""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length "
        + str(len(stream)).encode()
        + b" >>\nstream\n"
        + stream
        + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    trailer = f"<< /Size {len(objects) + 1} /Root 1 0 R >>"
    out += f"trailer\n{trailer}\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def test_excel_becomes_csv_per_sheet() -> None:
    got = extract_text(_xlsx(), XLSX)
    assert (
        got.text
        == "## Sheet: 利润表\n年份,营收\n2025,100\n\n## Sheet: 说明\n单位：万元\n\n"
    )
    assert not got.truncated and got.note is None


def test_pdf_text_comes_out_page_by_page() -> None:
    got = extract_text(_pdf("Revenue 100"), "application/pdf")
    assert got.text.startswith("## Page 1\nRevenue 100")
    assert got.note is None


def test_a_pdf_without_text_says_so() -> None:
    got = extract_text(_pdf(""), "application/pdf")
    assert got.note == NO_TEXT


def test_broken_files_are_unreadable_not_errors() -> None:
    assert extract_text(b"%PDF-1.4 broken", "application/pdf").note == UNREADABLE
    assert extract_text(b"PK\x03\x04junk", XLSX).note == UNREADABLE
    assert extract_text(b"x", "image/png").note == UNREADABLE


def test_a_workbook_that_unpacks_too_large_is_not_opened(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(text_extract, "MAX_XLSX_UNZIPPED", 10)
    assert extract_text(_xlsx(), XLSX).note == UNREADABLE


def test_text_is_cut_at_the_limit(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(text_extract, "MAX_TEXT_CHARS", 10)
    got = extract_text(("一" * 50).encode(), "text/plain; charset=utf-8")
    assert got.text == "一" * 10 and got.truncated and got.note


async def test_the_same_content_is_extracted_once(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    calls: list[str] = []

    def fake(data: bytes, content_type: str) -> Extracted:
        calls.append(content_type)
        return Extracted(text=data.decode())

    monkeypatch.setattr(text_extract, "extract_text", fake)
    extractor = CachedTextExtractor()
    sha = hashlib.sha256(b"abc").hexdigest()
    got = await asyncio.gather(
        *[
            extractor.extract(data=b"abc", content_type="text/plain", sha256=sha)
            for _ in range(3)
        ]
    )
    assert [g.text for g in got] == ["abc"] * 3 and calls == ["text/plain"]


async def test_a_reader_that_gives_up_still_leaves_the_text_cached(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """抽得慢、读的那头先超时走了：抽取照样做完进缓存，下一次直接拿到（K19 预演）。"""
    calls: list[str] = []
    gate = threading.Event()

    def slow(data: bytes, content_type: str) -> Extracted:
        calls.append(content_type)
        gate.wait(5)
        return Extracted(text=data.decode())

    monkeypatch.setattr(text_extract, "extract_text", slow)
    extractor = CachedTextExtractor()
    sha = hashlib.sha256(b"abc").hexdigest()
    with pytest.raises(TimeoutError):
        await asyncio.wait_for(
            extractor.extract(data=b"abc", content_type="application/pdf", sha256=sha),
            timeout=0.05,
        )
    gate.set()
    for _ in range(100):  # 等后台那次抽完
        if not extractor._pending:
            break
        await asyncio.sleep(0.01)
    got = await extractor.extract(
        data=b"abc", content_type="application/pdf", sha256=sha
    )
    assert got.text == "abc" and calls == ["application/pdf"]


# ---------- 按格子读 ----------


def _statements() -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "利润表"
    sheet.append(["科目", dt.date(2024, 3, 31), dt.datetime(2024, 6, 30), None])
    sheet.append(["营业收入", 46485000000.12, "1,234.50", None])
    sheet.append(["营业成本", "(3,000)", "−12", None])
    sheet.append([None, None, None, None])
    sheet.append(["其中：利息", 0.1, "-", None])
    sheet.append(["其中：利息", "12.5%", "", None])
    book.create_sheet("说明").append(["单位：元"])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def test_excel_becomes_labelled_cells_per_sheet() -> None:
    tables = extract_tables(_statements(), XLSX)
    sheet = tables.sheet("利润表")
    assert sheet is not None and tables.sheet(None) == sheet
    assert sheet.columns == ("2024-03-31", "2024-06-30")
    assert [r.label for r in sheet.rows] == [
        "营业收入",
        "营业成本",
        "其中：利息",
        "其中：利息#2",
    ]
    income = sheet.row("营业收入")
    assert income is not None
    assert income.cells["2024-03-31"].number == "46485000000.12"
    assert (income.cells["2024-06-30"].raw, income.cells["2024-06-30"].number) == (
        "1,234.50",
        "1234.5",
    )
    cost = sheet.row("营业成本")
    assert cost is not None
    assert [c.number for c in cost.cells.values()] == ["-3000", "-12"]
    interest = sheet.row("其中：利息")
    assert interest is not None and interest.cells["2024-03-31"].number == "0.1"
    assert interest.cells["2024-06-30"].number is None
    second = sheet.row("其中：利息#2")
    assert (
        second is not None and second.cells["2024-03-31"].number is None
    )  # 百分号不猜
    assert [s.name for s in tables.sheets] == ["利润表", "说明"]


def test_repeated_column_labels_are_numbered_not_overwritten() -> None:
    """合并与母公司并排、同一报告期出现两次：后一列不能盖掉前一列（审读第 5 条）。"""
    tables = extract_tables(
        "科目,2024-12-31,2024-12-31,,2023-12-31\n营收,100,60,x,90\n".encode(),
        "text/csv",
    )
    sheet = tables.sheet(None)
    assert sheet is not None
    assert sheet.columns == ("2024-12-31", "2024-12-31#2", "2023-12-31")
    row = sheet.row("营收")
    assert row is not None
    assert {k: c.number for k, c in row.cells.items()} == {
        "2024-12-31": "100",
        "2024-12-31#2": "60",
        "2023-12-31": "90",
    }


def test_xls_date_cells_become_dates_not_serial_numbers() -> None:
    import xlrd

    class XlsCell:
        def __init__(self, ctype: int, value: object) -> None:
            self.ctype, self.value = ctype, value

    assert _xls_value(XlsCell(xlrd.XL_CELL_DATE, 45657.0), 0) == dt.datetime(
        2024, 12, 31
    )
    assert _xls_value(XlsCell(xlrd.XL_CELL_NUMBER, 45657.0), 0) == 45657.0
    assert _xls_value(XlsCell(xlrd.XL_CELL_DATE, 1e20), 0) == 1e20  # 坏日期照原值
    header = [_xls_value(XlsCell(xlrd.XL_CELL_TEXT, "科目"), 0)] + [
        _xls_value(XlsCell(xlrd.XL_CELL_DATE, 45657.0), 0)
    ]
    sheet = _sheet("利润表", [header, ["营收", 100.0]])
    assert sheet.columns == ("2024-12-31",)


def test_csv_is_one_table_named_data() -> None:
    tables = extract_tables(
        "\ufeff科目,2024\n营收,100\n".encode(), "text/csv; charset=utf-8"
    )
    assert [(s.name, s.columns) for s in tables.sheets] == [("data", ("2024",))]


@pytest.mark.parametrize(
    ("value", "number"),
    [
        (100, "100"),
        (1.5, "1.5"),
        (float("nan"), None),
        (float("inf"), None),
        (True, None),
        ("1e400", None),
        ("0.00", "0"),
        ("-0", "0"),
        ("（12）", None),
        ("abc", None),
        ("--", None),
        (dt.date(2024, 1, 1), None),
    ],
)
def test_numbers_are_read_conservatively(value: object, number: str | None) -> None:
    assert number_of(value) == number


def test_non_tables_and_broken_files_are_unreadable() -> None:
    for data, kind in [
        (b"x", "text/plain"),
        (b"PK\x03\x04junk", XLSX),
        (b"junk", "application/vnd.ms-excel"),
        (b"%PDF-1.4", "application/pdf"),
    ]:
        with pytest.raises(TableUnreadable):
            extract_tables(data, kind)


def test_tables_have_size_limits(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(table_extract, "MAX_XLSX_UNZIPPED", 10)
    with pytest.raises(TableUnreadable):
        extract_tables(_statements(), XLSX)
    monkeypatch.setattr(table_extract, "MAX_ROWS", 1)
    with pytest.raises(TableUnreadable):
        extract_tables(b"k,v\na,1\nb,2\n", "text/csv")


def test_a_table_page_holds_about_four_thousand_cells() -> None:
    from app.domain.knowbots.data_read import Row

    rows = tuple(Row(label=str(i), cells={}) for i in range(450))
    sheet = Sheet(name="s", columns=tuple(str(i) for i in range(10)), rows=rows)
    first, pages = table_page(sheet, 1)
    assert (len(first), pages) == (400, 2)
    assert len(table_page(sheet, 2)[0]) == 50 and table_page(sheet, 3)[0] == ()
    assert table_page(Sheet(name="e", columns=(), rows=()), 1) == ((), 1)


def test_a_wide_table_page_stays_within_the_bytes_the_model_can_see() -> None:
    """格子不到 4000，但字节超了：按字节再分页（待办账 137）。"""
    from app.domain.knowbots.data_read import Cell, Row

    columns = tuple(f"2025-{m:02d}-30" for m in range(1, 13))
    rows = tuple(
        Row(
            label=f"归属于母公司所有者权益（或股东权益）合计{i}",
            cells={c: Cell(raw="170,899,152,276.34") for c in columns},
        )
        for i in range(300)
    )
    sheet = Sheet(name="资产负债表", columns=columns, rows=rows)
    _, pages = table_page(sheet, 1)
    assert pages > 1
    shown = [table_page(sheet, p)[0] for p in range(1, pages + 1)]
    assert sum(len(part) for part in shown) == 300
    for part in shown:
        body = {
            "columns": list(columns),
            "rows": [
                {"label": r.label, "values": {c: x.raw for c, x in r.cells.items()}}
                for r in part
            ],
        }
        assert len(json.dumps(body, ensure_ascii=False).encode("utf-8")) <= PAGE_BYTES


async def test_the_same_table_is_parsed_once(monkeypatch: pytest.MonkeyPatch) -> None:
    calls: list[str] = []
    real = table_extract.extract_tables

    def counting(data: bytes, content_type: str):  # type: ignore[no-untyped-def]
        calls.append(content_type)
        return real(data, content_type)

    monkeypatch.setattr(table_extract, "extract_tables", counting)
    extractor = CachedTableExtractor()
    data = b"k,v\na,1\n"
    sha = hashlib.sha256(data).hexdigest()
    for _ in range(3):
        await extractor.tables(data=data, content_type="text/csv", sha256=sha)
    assert calls == ["text/csv"]


def test_cell_requests_are_bounded() -> None:
    ref = {"dataset_id": KNOWBOT, "sheet": "s", "row": "r", "column": "c"}
    base = {
        "contract": "knowbot-data/v1",
        "purpose": {"kind": "test", "run_id": "knowlab-run-0002"},
    }
    with pytest.raises(ValueError):
        CellsRequestV1.model_validate(base | {"cells": []})
    with pytest.raises(ValueError):
        CellsRequestV1.model_validate(base | {"cells": [ref] * 501})
    with pytest.raises(ValueError):
        CellsRequestV1.model_validate(base | {"cells": [ref | {"row": ""}]})
    assert (
        len(CellsRequestV1.model_validate(base | {"cells": [ref] * 500}).cells) == 500
    )


# ---------- 核引文（SDD 0036） ----------

REPORT = (
    "## Page 1\n封面\n\n"
    "## Page 2\n第三节 管理层讨论与分析\n"
    "报告期内，公司营业收入同比增长\n15.7%，主要\n\n"
    "## Page 3\n得益于系列酒销量增加；毛利率略有下降。\n\n"
)


def test_pdf_pages_follow_the_extraction_markers() -> None:
    pages = pdf_pages(REPORT)
    assert sorted(pages) == [1, 2, 3]
    assert "管理层讨论与分析" in pages[2] and "封面" not in pages[2]
    assert pdf_pages("没有页的文字") == {}


def test_quotes_match_after_normalizing_spaces_and_widths() -> None:
    pages = pdf_pages(REPORT)
    assert (
        normalize_quote(" 营业收入 同比\n增长　１５．７％ ") == "营业收入同比增长15.7%"
    )
    assert find_quote(pages, 2, "公司营业收入同比增长 15.7%") is None
    assert find_quote(pages, 2, "公司营业收入同比增长１５．７％") is None  # 全角
    assert (
        find_quote(pages, 2, "公司营业收入同比增长 15.8%") == "not_found"
    )  # 改了一个字
    assert find_quote(pages, 3, "公司营业收入同比增长 15.7%") == "not_found"  # 页码不对


def test_a_quote_may_run_onto_the_next_page_but_must_start_on_its_own() -> None:
    pages = pdf_pages(REPORT)
    assert find_quote(pages, 2, "15.7%，主要得益于系列酒销量增加") is None
    # 整段都在第 3 页：写成第 2 页不算
    assert find_quote(pages, 2, "得益于系列酒销量增加；毛利率略有下降") == "not_found"
    assert find_quote(pages, 3, "得益于系列酒销量增加；毛利率略有下降") is None


def test_quote_misses_say_why() -> None:
    pages = pdf_pages(REPORT)
    assert find_quote(pages, 2, "营业收入") == "too_short"
    assert find_quote(pages, 2, "字" * 1001) == "too_long"
    assert find_quote({}, 1, "公司营业收入同比增长") == "no_pages"
    assert find_quote(pages, 9, "公司营业收入同比增长") == "page_out_of_range"


def test_spaces_between_letters_and_digits_still_count() -> None:
    """去掉所有空白会让「now here」等于「nowhere」、「1 234」等于「12 34」。

    审读 10-19。
    """
    pages = {1: "This is now here in the report. Annual sales: 1 234 million yuan."}
    assert find_quote(pages, 1, "This is now here in the report.") is None
    assert find_quote(pages, 1, "This is nowhere in the report.") == "not_found"
    assert find_quote(pages, 1, "Annual sales: 12 34 million") == "not_found"
    assert (
        find_quote(pages, 1, "Annual  sales:\n1 234 million") is None
    )  # 断行、多空格照认


def test_numbers_alone_are_too_short_to_prove_a_passage() -> None:
    pages = {1: "截至2024年12月31日，营业收入为1,234,567.89万元"}
    assert find_quote(pages, 1, "2024年12月31日") == "too_short"
    assert find_quote(pages, 1, "1,234,567.89万元") == "too_short"
    assert (
        find_quote(pages, 1, "截至2024年12月31日，营业收入为1,234,567.89万元") is None
    )


def test_a_quote_past_a_truncated_last_page_is_out_of_range() -> None:
    pages = {1: "第一页的叙述性文字一段", 2: "被截断的最后一页只抽到了前面"}
    quote = "最后一页后面没抽到的那一段原话"
    assert find_quote(pages, 2, quote) == "not_found"
    assert find_quote(pages, 2, quote, truncated=True) == "page_out_of_range"
    assert find_quote(pages, 1, quote, truncated=True) == "not_found"


def test_quote_requests_are_bounded() -> None:
    quote = {"dataset_id": KNOWBOT, "page": 1, "text": "公司营业收入同比增长"}
    base = {
        "contract": "knowbot-data/v1",
        "purpose": {"kind": "test", "run_id": "knowlab-run-0002"},
    }
    with pytest.raises(ValueError):
        QuotesRequestV1.model_validate(base | {"quotes": []})
    with pytest.raises(ValueError):
        QuotesRequestV1.model_validate(base | {"quotes": [quote] * 101})
    with pytest.raises(ValueError):
        QuotesRequestV1.model_validate(base | {"quotes": [quote | {"page": 0}]})
    assert len(QuotesRequestV1.model_validate(base | {"quotes": [quote] * 100}).quotes)


def _quotes(purpose: dict[str, str], *refs: tuple[str, int, str]) -> QuotesRequestV1:
    return QuotesRequestV1.model_validate(
        {
            "contract": "knowbot-data/v1",
            "purpose": purpose,
            "quotes": [{"dataset_id": d, "page": p, "text": t} for d, p, t in refs],
        }
    )


async def _bind(world: World, record: DatasetRecord) -> None:
    """把这份数据也绑进在测试的那一版。"""
    doc = dict(world.repo.drafts[world.testing.id].manifest)
    doc["data_bindings"] = doc["data_bindings"] + [
        {"dataset_id": str(record.id), "version": "1", "sha256": record.sha256}
    ]
    await world.repo.update_draft(world.testing.id, manifest=doc)


async def test_quotes_are_checked_in_the_bound_pdf(world: World) -> None:
    report = await world._dataset(
        "2025年报.pdf", "application/pdf", _pdf("Revenue grew because more was sold")
    )
    await _bind(world, report)
    purpose = {"kind": "test", "run_id": "knowlab-run-0002"}
    checked = await world.reader.quotes(
        "knowbench",
        _quotes(
            purpose,
            (str(report.id), 1, "Revenue grew because more was sold"),
            (str(report.id), 1, "Revenue fell because less was sold"),
            (str(report.id), 2, "Revenue grew because more was sold"),
            (str(world.bound.id), 1, "公司,营收 A,100 不是 PDF"),
        ),
    )
    assert [(q.found, q.reason) for q in checked.quotes] == [
        (True, None),
        (False, "not_found"),
        (False, "page_out_of_range"),
        (False, "no_pages"),
    ]
    assert checked.quotes[0].sha256 == report.sha256
    assert checked.quotes[0].version == "1"


async def test_quotes_keep_the_same_two_checks(world: World) -> None:
    purpose = {"kind": "test", "run_id": "knowlab-run-0002"}
    request = _quotes(purpose, (str(world.loose.id), 1, "不该读到的那份数据里的字"))
    with pytest.raises(NotFoundError):
        await world.reader.quotes("knowbench", request)  # 没绑在这一版
    bound = _quotes(purpose, (str(world.bound.id), 1, "公司营收一百万元整的那句"))
    with pytest.raises(NotFoundError):
        await world.reader.quotes("other-runtime", bound)
    with pytest.raises(NotFoundError):
        await world.reader.quotes(
            "knowbench",
            _quotes(
                {"kind": "test", "run_id": "knowlab-run-9999"},
                (str(world.bound.id), 1, "公司营收一百万元整的那句"),
            ),
        )


# ---------- 契约文件 ----------


@pytest.mark.parametrize("name", sorted(DOCUMENTS))
def test_schema_files_are_generated_from_the_models(name: str) -> None:
    if not CONTRACT_DIR.exists():
        pytest.skip("parent contracts/ is not checked out")
    assert (CONTRACT_DIR / name).read_text(encoding="utf-8") == render(name), (
        "regenerate: python -m app.domain.knowbots.data_schema"
    )


def test_contract_manifest_digests_match_the_schemas() -> None:
    if not CONTRACT_DIR.exists():
        pytest.skip("parent contracts/ is not checked out")
    release = json.loads((CONTRACT_DIR / "contract-manifest.json").read_text())
    assert (release["contract"], release["major_version"]) == ("knowbot-data", 1)
    assert set(release["schemas"]) == set(DOCUMENTS)
    for name, digest in release["schemas"].items():
        assert hashlib.sha256((CONTRACT_DIR / name).read_bytes()).hexdigest() == digest


# ---------- 内部接口 ----------


def _client(base_url: str = "http://127.0.0.1") -> httpx.AsyncClient:
    # 内部接口只认集群内与回环地址（审读第 1 条）
    return httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url=base_url)


async def test_public_hosts_do_not_see_the_internal_entry() -> None:
    body = {"contract": "knowbot-data/v1", "purpose": {"kind": "test"}}
    async with _client("http://testserver") as h:
        hidden = await h.post(
            f"{INTERNAL}/list", json=body, headers={"Authorization": "Bearer x"}
        )
    assert hidden.status_code == 404


async def test_without_the_relation_configured_everything_is_refused(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    for key in list(os.environ):
        if key.startswith("SERVICE_AUTH_READ"):
            monkeypatch.delenv(key)
    body = {
        "contract": "knowbot-data/v1",
        "purpose": {"kind": "test", "run_id": "knowlab-run-0002"},
    }
    async with _client() as h:
        assert (await h.post(f"{INTERNAL}/list", json=body)).status_code == 401
        refused = await h.post(
            f"{INTERNAL}/list", json=body, headers={"Authorization": "Bearer x"}
        )
    assert refused.status_code == 503
    assert refused.json()["code"] == "service_relation_unconfigured"


async def test_the_runtime_lists_and_reads_through_the_internal_entry(
    world: World,
) -> None:
    app.dependency_overrides[require_read_service] = lambda: ServiceCaller(
        caller="knowbench",
        principal=None,  # type: ignore[arg-type]
    )
    app.dependency_overrides[knowbot_data_routes.reader] = lambda: world.reader
    try:
        purpose = world.knowbot()
        async with _client() as h:
            listed = await h.post(
                f"{INTERNAL}/list",
                json={"contract": "knowbot-data/v1", "purpose": purpose},
            )
            read = await h.post(
                f"{INTERNAL}/read",
                json={
                    "contract": "knowbot-data/v1",
                    "purpose": purpose,
                    "dataset_id": str(world.bound.id),
                },
            )
            other = await h.post(
                f"{INTERNAL}/read",
                json={
                    "contract": "knowbot-data/v1",
                    "purpose": purpose,
                    "dataset_id": str(world.loose.id),
                },
            )
            table = await h.post(
                f"{INTERNAL}/table",
                json={
                    "contract": "knowbot-data/v1",
                    "purpose": purpose,
                    "dataset_id": str(world.bound.id),
                },
            )
            cells = await h.post(
                f"{INTERNAL}/cells",
                json={
                    "contract": "knowbot-data/v1",
                    "purpose": purpose,
                    "cells": [
                        {
                            "dataset_id": str(world.bound.id),
                            "sheet": "data",
                            "row": "A",
                            "column": "营收",
                        }
                    ],
                },
            )
            bad = await h.post(f"{INTERNAL}/list", json={"purpose": purpose})
    finally:
        app.dependency_overrides.clear()
    assert listed.status_code == 200 and listed.headers["cache-control"] == "no-store"
    assert [d["name"] for d in listed.json()["datasets"]] == ["年报.csv"]
    assert read.json()["text"] == "公司,营收\nA,100\n"
    assert other.status_code == 404 and other.json()["code"] == "knowbot_data_not_found"
    assert table.json()["rows"] == [{"label": "A", "values": {"营收": "100"}}]
    assert cells.headers["cache-control"] == "no-store"
    assert cells.json()["cells"][0]["number"] == "100"
    assert bad.status_code == 422


async def test_the_runtime_checks_quotes_through_the_internal_entry(
    world: World,
) -> None:
    app.dependency_overrides[require_read_service] = lambda: ServiceCaller(
        caller="knowbench",
        principal=None,  # type: ignore[arg-type]
    )
    app.dependency_overrides[knowbot_data_routes.reader] = lambda: world.reader
    report = await world._dataset(
        "2025年报.pdf", "application/pdf", _pdf("Revenue grew because more was sold")
    )
    doc = dict(world.repo.drafts[world.delivered.id].manifest)
    doc["data_bindings"] = doc["data_bindings"] + [
        {"dataset_id": str(report.id), "version": "1", "sha256": report.sha256}
    ]
    await world.repo.update_draft(world.delivered.id, manifest=doc)
    try:
        body = {
            "contract": "knowbot-data/v1",
            "purpose": world.knowbot(),
            "quotes": [
                {"dataset_id": str(report.id), "page": 1, "text": "more was sold"},
                {"dataset_id": str(report.id), "page": 1, "text": "less was sold"},
            ],
        }
        async with _client() as h:
            checked = await h.post(f"{INTERNAL}/quotes", json=body)
            bad = await h.post(f"{INTERNAL}/quotes", json=body | {"quotes": []})
    finally:
        app.dependency_overrides.clear()
    assert checked.status_code == 200
    assert checked.headers["cache-control"] == "no-store"
    assert [(q["found"], q["reason"]) for q in checked.json()["quotes"]] == [
        (True, None),
        (False, "not_found"),
    ]
    assert bad.status_code == 422
